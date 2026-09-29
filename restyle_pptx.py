#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Restyle the Persian master's thesis-defense PowerPoint template.

Hard constraints honored (verified after the run):
  * every slide, shape, position, size and text string is preserved exactly
  * right-side section navigation stays, with the active-section indication
  * slide numbers (bottom-right of the sidebar) stay untouched
  * 16:9 widescreen (12191695 x 6858000 EMU) preserved
  * Persian paragraphs remain RTL (pPr rtl="1" untouched)

Visual restyle applied:
  * restrained technical palette: deep navy, white, light gray, teal accent
  * right-side nav rendered as smooth pill/capsule shapes (full rounding)
  * soft rounded corners on panels and bars
  * subtle soft outer shadows (sidebar, nav pills, content panel) - no more
  * typography: B Nazanin for Persian (complex script), Times New Roman for
    Latin / English technical terms
  * slide titles 28 pt (spec 26-30), body text kept 18 pt (spec 18-20)
"""

from lxml import etree
from pptx import Presentation

SRC = "قالب_دفاع_پایان_نامه_50_اسلایدی_v2.pptx"
DST = "قالب_دفاع_پایان_نامه_50_اسلایدی_v3.pptx"

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def qn(tag):
    pre, local = tag.split(":")
    ns = {"a": A, "p": P}[pre]
    return "{%s}%s" % (ns, local)


# ----------------------------------------------------------------------------
# 1. Palette: refined deep navy / light gray / teal (subtle, academic)
# ----------------------------------------------------------------------------
COLOR_MAP = {
    "173B57": "132F4C",   # deep navy  (titles, active pill, cover bar, numbers)
    "2C7A8A": "2A7E94",   # teal accent (bars, active indicator, key note)
    "1B2833": "22303F",   # primary dark text (inactive nav labels)
    "7A8793": "6E7F8E",   # muted text (captions, placeholders, footer)
    "F2F4F7": "EDF1F6",   # light gray surfaces (sidebar, cover band)
    "F7F8FA": "F8FAFC",   # inactive nav pills (near-white on gray sidebar)
    "DCE2E7": "D4DEE7",   # dividers
    "FBFCFD": "FCFDFE",   # content panel background
    "E2E7EB": "DDE5ED",   # content panel border
    # FFFFFF stays pure white
}

# ----------------------------------------------------------------------------
# 2. Corner radii (a:avLst adj value -> new value).  roundRect adj is the
#    corner radius as a fraction of the shorter side; 50000 = full capsule.
# ----------------------------------------------------------------------------
ADJ_MAP = {
    "26667": "50000",   # nav section pills        -> smooth full pill
    "42857": "50000",   # top accent bar           -> capsule
    "11023": "17000",   # sidebar panel            -> smoother curved panel
    "45455": "50000",   # sidebar accent strip     -> capsule
    "44444": "50000",   # active-section indicator -> capsule
    "55556": "50000",   # sidebar divider          -> capsule
    "50000": "50000",   # title underline          -> keep capsule
    "3048":  "6500",    # content panel            -> soft rounded corners
    "27273": "50000",   # cover top bar            -> capsule
    "40000": "50000",   # cover accent line        -> capsule
    "5769":  "15000",   # cover bottom band        -> soft rounded corners
}

# Shape sizes (EMU) used to identify roles for shadows - the template uses a
# single consistent geometry set across all 50 slides.
W_SIDEBAR = 1493215
W_NAVPILL = 1337767
H_NAVPILL = 548640
H_CONTENT = 4800600

# Shadow recipe: (blurRad, dist, dir, alpha)   dir: 2700000=45deg down-right,
# 8100000=135deg down-left (used for the right-edge sidebar).
SHADOW_DEFAULT = (76200, 25400, 2700000, 11000)        # content panel
SHADOW_NAV_INACTIVE = (63500, 22860, 2700000, 9000)    # inactive pills
SHADOW_NAV_ACTIVE = (88900, 31750, 2700000, 16000)     # active pill (a bit more)
SHADOW_SIDEBAR = (101600, 38100, 8100000, 13000)       # sidebar panel


def remap_colors(root):
    for el in root.iter(qn("a:srgbClr")):
        val = el.get("val")
        if val in COLOR_MAP:
            el.set("val", COLOR_MAP[val])


def remap_adj(root):
    for geom in root.iter(qn("a:prstGeom")):
        if geom.get("prst") != "roundRect":
            continue
        av = geom.find(qn("a:avLst"))
        if av is None or len(av) == 0:
            continue
        gd = av[0]
        fmla = gd.get("fmla") or ""
        if fmla.startswith("val "):
            old = fmla[4:].strip()
            if old in ADJ_MAP:
                gd.set("fmla", "val %s" % ADJ_MAP[old])


def add_shadow(spPr, blur, dist, direction, alpha):
    for el in spPr.findall(qn("a:effectLst")):
        spPr.remove(el)
    xml = (
        '<a:effectLst xmlns:a="%s">'
        '<a:outerShdw blurRad="%d" dist="%d" dir="%d" rotWithShape="0">'
        '<a:srgbClr val="132F4C"><a:alpha val="%d"/></a:srgbClr>'
        "</a:outerShdw></a:effectLst>" % (A, blur, dist, direction, alpha)
    )
    eff = etree.fromstring(xml)
    ln = spPr.find(qn("a:ln"))
    if ln is not None:
        ln.addnext(eff)          # schema order: ... fill, ln, effectLst ...
    else:
        spPr.append(eff)


def style_shapes(slide_el):
    """Apply shadows per shape role (identified by geometry, unchanged layout)."""
    for sp in slide_el.iter(qn("p:sp")):
        spPr = sp.find(qn("p:spPr"))
        if spPr is None:
            continue
        xfrm = spPr.find(qn("a:xfrm"))
        if xfrm is None:
            continue
        ext = xfrm.find(qn("a:ext"))
        if ext is None:
            continue
        w, h = int(ext.get("cx")), int(ext.get("cy"))
        if w == W_SIDEBAR and h > 6_000_000:            # sidebar panel
            add_shadow(spPr, *SHADOW_SIDEBAR)
        elif w == W_NAVPILL and h == H_NAVPILL:          # nav section pills
            # active pill = solidFill navy(132F4C after remap); inactive = F8FAFC
            fill = spPr.find(qn("a:solidFill"))
            clr = fill[0].get("val") if fill is not None and len(fill) else ""
            if clr == "132F4C":
                add_shadow(spPr, *SHADOW_NAV_ACTIVE)
            else:
                add_shadow(spPr, *SHADOW_NAV_INACTIVE)
        elif h == H_CONTENT and w > 8_000_000:           # content panel
            add_shadow(spPr, *SHADOW_DEFAULT)


def set_run_typefaces(root):
    """Persian (complex script) -> B Nazanin, Latin -> Times New Roman."""
    faces = (("a:latin", "Times New Roman"),
             ("a:ea", "Times New Roman"),
             ("a:cs", "B Nazanin"))
    trailing = (qn("a:sym"), qn("a:hlinkClick"), qn("a:hlinkMouseOver"),
                qn("a:rtl"), qn("a:extLst"))
    for rPr in root.iter():
        if rPr.tag not in (qn("a:rPr"), qn("a:endParaRPr")):
            continue
        # title size bump 26 pt -> 28 pt (spec range 26-30 pt)
        if rPr.get("sz") == "2600":
            rPr.set("sz", "2800")
        # drop old font elements and re-add in schema order: latin, ea, cs
        # (must sit after fill/effect and before sym/hlink/rtl/extLst)
        for tag, _ in faces:
            for el in rPr.findall(qn(tag)):
                rPr.remove(el)
        anchor = None
        for child in rPr:
            if child.tag in trailing:
                anchor = child
                break
        for tag, face in faces:
            el = etree.Element(qn(tag))
            el.set("typeface", face)
            if anchor is not None:
                anchor.addprevious(el)
            else:
                rPr.append(el)


def patch_theme_fonts(prs):
    """Set theme font scheme so freshly typed text follows the same rule."""
    for part in prs.part.package.iter_parts():
        if not str(part.content_type).endswith("theme+xml"):
            continue
        try:
            root = etree.fromstring(part.blob)
        except Exception:
            continue
        for coll_tag in ("a:majorFont", "a:minorFont"):
            coll = root.find(".//" + qn(coll_tag))
            if coll is None:
                continue
            for tag, face in (("a:latin", "Times New Roman"),
                              ("a:ea", "Times New Roman"),
                              ("a:cs", "B Nazanin")):
                el = coll.find(qn(tag))
                if el is not None:
                    el.set("typeface", face)
            for font_el in coll.findall(qn("a:font")):
                script = font_el.get("script", "")
                if script == "Arab":
                    font_el.set("typeface", "B Nazanin")
                elif script == "Latn":
                    font_el.set("typeface", "Times New Roman")
        part._blob = etree.tostring(root, xml_declaration=True,
                                    encoding="UTF-8", standalone=True)


def main():
    prs = Presentation(SRC)

    for slide in prs.slides:
        root = slide._element
        remap_colors(root)
        remap_adj(root)
        style_shapes(root)          # before? after? fills already remapped -> use new navy
        set_run_typefaces(root)

    patch_theme_fonts(prs)
    prs.save(DST)
    print("saved:", DST)


if __name__ == "__main__":
    main()
